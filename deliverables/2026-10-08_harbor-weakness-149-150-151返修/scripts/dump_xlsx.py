"""Dump every sheet/cell of an xlsx to text for reading."""
import sys
from openpyxl import load_workbook

path = sys.argv[1]
wb = load_workbook(path, data_only=True)
for ws in wb.worksheets:
    print(f"===== SHEET: {ws.title}  dims={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column} =====")
    for row in ws.iter_rows():
        vals = []
        for c in row:
            v = c.value
            if v is None:
                continue
            vals.append(f"[{c.coordinate}] {v}")
        if vals:
            print(" | ".join(vals))
    print()
