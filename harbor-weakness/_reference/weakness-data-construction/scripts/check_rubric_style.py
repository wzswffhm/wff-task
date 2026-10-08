#!/usr/bin/env python3
"""判据风格与锚点自检（人检高频退回项）。

用法:
    python check_rubric_style.py <task-dir> [...] [--strict]

检查：提问式判据、无定义量词、表格题定位语混入、criterion_type 误标，
以及判据锚点能否在 environment/input_files 里溯源的待人工核实清单。
默认只提示（退出码 0）；--strict 时把提问式/量词/措辞三项计入 FAIL。

溯源部分尽力而为：能装到 pdfplumber / python-docx / openpyxl 就提取材料正文比对，
装不到就只输出待核实清单，不做判定。
"""
import json
import os
import re
import sys

VAGUE = ("若干", "个别", "普遍", "多数", "一些", "少量", "大致", "大约", "酌情")
VAGUE_WHITELIST = ("全部或部分", "部分使用", "独立部分")   # 规则定义原文的说法，不算模糊量词
VAGUE_COUNTED = re.compile(r"[一二三四五六七八九十两0-9]+\s*个?\s*部分")   # 「两部分」这类量词，不是模糊量词
SHEET_WORDS = ("单元格", "公式", "逐格复算", "可逐格")
ANCHOR_QUOTED = re.compile(r"[\u300c\u300e\"\u201c]([^\u300d\u300f\"\u201d]{4,80})[\u300d\u300f\"\u201d]")
# 值得溯源的"自造口径"短语：XX口径 / XX阈值 / XX限额，或纯数字门槛
ANCHOR_NEAR = re.compile(r"([\u4e00-\u9fff]{2,12}(?:口径|阈值|限额))")
ANCHOR_NUMBER = re.compile(
    r"(?:§\s*\d+(?:\.\d+)*|\b\d{3}\.\d{2,3}\b|10\^?\d{1,3}|\b\d+(?:\.\d+)?%|\b\d{1,3}(?:,\d{3})+\b)")
QUOTED_SKIP = ("依据", "结论", "问题", "出处", "判断把握度", "依据出处", "事项", "时点或条件",
               "材料冲突与口径说明", "存疑事项与风险提示", "交易推进的前置条件与时间表")
QUESTION = re.compile(r"^[^\uff0c\u3002\uff1b\uff1a]{0,8}?(是否|能否|有无)")
LEVEL_KEYS = ("1", "0.75", "0.5", "0.25", "0")


def load_items(task_dir):
    with open(os.path.join(task_dir, "rubrics.json"), encoding="utf-8") as fh:
        raw = json.load(fh)
    return raw["items"] if isinstance(raw, dict) else raw


def input_text(task_dir, limit_chars=400000):
    """尽力提取 input_files 的文本（缺库则返回 None）。"""
    src = os.path.join(task_dir, "environment", "input_files")
    if not os.path.isdir(src):
        return None
    chunks = []
    for name in sorted(os.listdir(src)):
        path = os.path.join(src, name)
        low = name.lower()
        try:
            if low.endswith(".pdf"):
                import pdfplumber  # type: ignore

                with pdfplumber.open(path) as pdf:
                    chunks.extend((p.extract_text() or "") for p in pdf.pages)
            elif low.endswith(".docx"):
                from docx import Document  # type: ignore

                doc = Document(path)
                chunks.extend(p.text for p in doc.paragraphs)
                for table in doc.tables:
                    for row in table.rows:
                        chunks.append(" ".join(c.text for c in row.cells))
            elif low.endswith((".xlsx", ".xlsm")):
                import openpyxl  # type: ignore

                wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
                for sheet in wb.worksheets:
                    for row in sheet.iter_rows(values_only=True):
                        chunks.append(" ".join("" if c is None else str(c) for c in row))
            else:
                continue
        except Exception:
            continue
    text = "\n".join(chunks)
    return text[:limit_chars] if text else None


def style_check(task_dir, strict):
    items = load_items(task_dir)
    materials = input_text(task_dir)
    fails, notes = [], []
    print(f"===== {os.path.basename(task_dir)}（{len(items)} 条）")
    for it in items:
        rid = it["id"]
        desc = str(it.get("description") or "")
        levels = it.get("levels") or {}
        blob = desc + " " + " ".join(str(levels.get(k, "")) for k in LEVEL_KEYS)
        # 判据里的引号常包着表头名（如"产品形态 × 是否构成…"），先去掉再判提问式
        m = QUESTION.search(ANCHOR_QUOTED.sub("", desc))
        if m:
            (fails if strict else notes).append(f"{rid} 提问式表述：{m.group(0)!r}")
        scrubbed = VAGUE_COUNTED.sub("", blob)
        for safe in VAGUE_WHITELIST:
            scrubbed = scrubbed.replace(safe, "")
        hits = [w for w in VAGUE if w in scrubbed]
        if hits:
            (fails if strict else notes).append(f"{rid} 无定义量词 {hits}，请改量化区间")
        hit_words = [w for w in SHEET_WORDS if w in blob]
        if hit_words and not str(it.get("dimension", "")).startswith("内容质量-数值与计算"):
            (fails if strict else notes).append(f"{rid} 含表格类定位语 {hit_words}，确认是否贴切")
        if str(it.get("criterion_type")) == "Subjective":
            notes.append(f"{rid} 标为 Subjective，确认能否按可核验锚点改 Objective")
        anchors = set()
        for quoted in ANCHOR_QUOTED.findall(desc):
            if any(k in quoted for k in ("口径", "阈值", "限额", "标准")) \
                    and not any(s in quoted for s in QUOTED_SKIP):
                anchors.add(quoted)
        for a in sorted(anchors):
            if materials is None:
                notes.append(f"{rid} 锚点待人工核实：{a}")
            elif a not in materials:
                notes.append(f"{rid} 锚点 {a} 未在 input_files 命中，请确认出处"
                             f"（材料为英文原文时中文措辞不会命中，需人工判断）")
        if materials is not None:
            squashed = materials.replace(" ", "").replace("\n", "")
            for num in sorted(set(ANCHOR_NUMBER.findall(desc))):
                token = num.strip()
                if token and token.replace(" ", "").replace(",", "") not in squashed:
                    notes.append(f"{rid} 条款/数值 {token} 未在 input_files 命中，请人工确认")
    for msg in fails:
        print("  FAIL " + msg)
    for msg in notes:
        print("  NOTE " + msg)
    if not fails and not notes:
        print("  OK   未发现风格问题")
    print(f"  —— 提问式/量词/措辞问题 {len(fails)} 条，待确认 {len(notes)} 条"
          f"（材料文本比对：{'已启用' if materials is not None else '未启用'}）")
    return fails, notes


def main():
    args = sys.argv[1:]
    strict = "--strict" in args
    dirs = [a for a in args if not a.startswith("-")]
    if not dirs:
        raise SystemExit(__doc__)
    total_fail = 0
    for task_dir in dirs:
        fails, _notes = style_check(task_dir.rstrip("/\\"), strict)
        total_fail += len(fails)
    print(f"\nFAIL 合计: {total_fail}"
          + ("" if strict else "（当前为提示模式；加 --strict 可把风格项计为 FAIL）"))
    return 1 if total_fail else 0


if __name__ == "__main__":
    sys.exit(main())
