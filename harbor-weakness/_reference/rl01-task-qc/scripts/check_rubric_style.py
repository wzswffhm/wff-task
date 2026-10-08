#!/usr/bin/env python3
"""判据风格与锚点自检（人检高频退回项）。

用法:
    python check_rubric_style.py <task-dir> [...] [--strict]

检查：提问式判据、无定义量词、表格题定位语混入、criterion_type 误标，
以及判据锚点能否在 environment/input_files 里溯源的待人工核实清单。
默认只提示（退出码 0）；--strict 时把提问式/量词/措辞三项计入 FAIL。
--no-material 跳过 environment/input_files 的正文提取（大批量质检提速用，
此时锚点溯源只输出待人工核实清单，不做材料比对）。这是本 skill 相对上游脚本的本地增量。

溯源部分尽力而为：能装到 pdfplumber / python-docx / openpyxl 就提取材料正文比对，
装不到就只输出待核实清单，不做判定。
"""
import json
import os
import re
import sys

VAGUE = ("若干", "个别", "普遍", "多数", "一些", "少量", "大致", "大约", "酌情")
VAGUE_WHITELIST = ("全部或部分", "部分使用", "独立部分")   # 规则定义原文的说法，不算模糊量词
VAGUE_COUNTED = re.compile(r"[一二三四五六七八九十两0-9]+\s*个?\s*部分")   # 「两部分」这类量词，不是模糊量词
# 本地增量（2026-10-02）：法规正式名称里的「若干」不是无定义量词。
#   真实误报：LAW-011 R03「引用2022年2月23日修订的《关于审理非法集资刑事案件具体应用法律
#   若干问题的解释》」、R14「《关于办理洗钱刑事案件适用法律若干问题的解释》第二条」被判
#   「无定义量词 ['若干']」。判据引的是法规全名，必须原样保留。
VAGUE_LAW_TITLE = re.compile(r"若干(?:问题)?的(?:解释|规定|意见|通知|批复|答复|办法|规则)")
# 本地增量（2026-10-02）：法律-税务题里的「计税公式 / 计算公式」是分析对象本身，不是
#   表格题模板用语。真实误报：LAW-012（境外股票期权）R36「自行给出股份转让环节的具体
#   计税公式或具体比例税率」被判「含表格类定位语 ['公式']」。
FORMULA_IN_CONTEXT = re.compile(r"(?:计税|计算|换算|结算|估值|定价|分摊)公式")
SHEET_WORDS = ("单元格", "公式", "逐格复算", "可逐格")
# 本地增量：只有「逐格复算 / 可逐格」是强信号（表格题模板用语）；
# 「单元格 / 公式」在金融、投行、代码类题里是正常表述（真实案例：FIN1-SKILL-DEP-001 的
# EV/EBITDA 计算链写「EBITDA 公式」「单元格」被判 FAIL，属误报）。
STRONG_SHEET = ("逐格复算", "可逐格")
ANCHOR_QUOTED = re.compile(r"[\u300c\u300e\"\u201c]([^\u300d\u300f\"\u201d]{4,80})[\u300d\u300f\"\u201d]")
# 值得溯源的"自造口径"短语：XX口径 / XX阈值 / XX限额，或纯数字门槛
ANCHOR_NEAR = re.compile(r"([\u4e00-\u9fff]{2,12}(?:口径|阈值|限额))")
ANCHOR_NUMBER = re.compile(
    r"(?:§\s*\d+(?:\.\d+)*|\b\d{3}\.\d{2,3}\b|10\^?\d{1,3}|\b\d+(?:\.\d+)?%|\b\d{1,3}(?:,\d{3})+\b)")
QUOTED_SKIP = ("依据", "结论", "问题", "出处", "判断把握度", "依据出处", "事项", "时点或条件",
               "材料冲突与口径说明", "存疑事项与风险提示", "交易推进的前置条件与时间表")
QUESTION = re.compile(r"^[^\uff0c\u3002\uff1b\uff1a]{0,8}?(是否|能否|有无)")
LEVEL_KEYS = ("1", "0.75", "0.5", "0.25", "0")

# 本地增量（2026-10-01，甲方返修意见 #2 的机器化）：
#   DUP_SCALE   description 里重复给出 0–1 档位段。档位唯一来源是 levels；
#               description 再写一遍会在 toml 里渲染成「1–5 标度 + 0–1 标度」两套，
#               判官输出小数被 schema 拒或随机漂移（真实案例：LAW-002 R04/R05/R06/
#               R07/R10/R12/R20/R21/R22，LAW-004 R04/R05/R06/R10/R19/R20）。
#   META_PATCH  判据里替题面歧义兜底的解释语。出现说明题面本身有歧义没修，
#               用判据把「不影响判定」写死（真实案例：LAW-004 R12「…是本任务明确
#               要求的组成部分，不影响上述四部分的齐全性与顺序判定」）。
#               根因在题面，应回题面改，判据里不该出现。
DUP_SCALE = re.compile(r"按[^。；]{0,12}评分[:：]|0\.75\s*=|0\.25\s*=")
META_PATCH = re.compile(
    r"(不影响[^。；]{0,12}判定|不计入[^。；]{0,12}判定|本任务明确要求的组成部分"
    r"|不影响上述|不计入上述)")

# 本地增量：--no-material 时跳过材料正文提取
SKIP_MATERIAL = False


def load_items(task_dir):
    with open(os.path.join(task_dir, "rubrics.json"), encoding="utf-8") as fh:
        raw = json.load(fh)
    if isinstance(raw, list):
        return raw
    for key in ("items", "rubrics", "criteria"):
        if isinstance(raw.get(key), list):
            return raw[key]
    for value in raw.values():
        if isinstance(value, list):
            return value
    raise SystemExit("rubrics.json 顶层找不到条目列表字段")


def task_domain(task_dir):
    """读 task.toml 的 domain：法律题按严口径，其他领域放宽表格类定位语。"""
    path = os.path.join(task_dir, "task.toml")
    if not os.path.isfile(path):
        return ""
    m = re.search(r'(?m)^domain\s*=\s*"([^"]+)"', open(path, encoding="utf-8").read())
    return m.group(1) if m else ""


def input_text(task_dir, limit_chars=400000):
    """尽力提取 input_files 的文本（缺库则返回 None）。"""
    src = os.path.join(task_dir, "environment", "input_files")
    if not os.path.isdir(src):
        return None
    chunks = []
    for name in sorted(os.listdir(src)):
        path = os.path.join(src, name)
        low = name.lower()
        try:
            if low.endswith(".pdf"):
                import pdfplumber  # type: ignore

                with pdfplumber.open(path) as pdf:
                    chunks.extend((p.extract_text() or "") for p in pdf.pages)
            elif low.endswith(".docx"):
                from docx import Document  # type: ignore

                doc = Document(path)
                chunks.extend(p.text for p in doc.paragraphs)
                for table in doc.tables:
                    for row in table.rows:
                        chunks.append(" ".join(c.text for c in row.cells))
            elif low.endswith((".xlsx", ".xlsm")):
                import openpyxl  # type: ignore

                wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
                for sheet in wb.worksheets:
                    for row in sheet.iter_rows(values_only=True):
                        chunks.append(" ".join("" if c is None else str(c) for c in row))
            else:
                continue
        except Exception:
            continue
    text = "\n".join(chunks)
    return text[:limit_chars] if text else None


def style_check(task_dir, strict):
    items = load_items(task_dir)
    materials = None if SKIP_MATERIAL else input_text(task_dir)
    domain = task_domain(task_dir)
    strict_sheet = ("法律" in domain) or not domain
    fails, notes = [], []
    print(f"===== {os.path.basename(task_dir)}（{len(items)} 条）")
    for it in items:
        rid = it["id"]
        desc = str(it.get("description") or "")
        levels = it.get("levels") or {}
        blob = desc + " " + " ".join(str(levels.get(k, "")) for k in LEVEL_KEYS)
        # 判据里的引号常包着表头名（如"产品形态 × 是否构成…"），先去掉再判提问式
        m = QUESTION.search(ANCHOR_QUOTED.sub("", desc))
        if m:
            (fails if strict else notes).append(f"{rid} 提问式表述：{m.group(0)!r}")
        scrubbed = VAGUE_COUNTED.sub("", blob)
        scrubbed = VAGUE_LAW_TITLE.sub("", scrubbed)
        for safe in VAGUE_WHITELIST:
            scrubbed = scrubbed.replace(safe, "")
        hits = [w for w in VAGUE if w in scrubbed]
        if hits:
            (fails if strict else notes).append(f"{rid} 无定义量词 {hits}，请改量化区间")
        hit_words = [w for w in SHEET_WORDS if w in FORMULA_IN_CONTEXT.sub("", blob)]
        if hit_words and not strict_sheet:
            hit_words = [w for w in hit_words if w in STRONG_SHEET]
        if hit_words and not str(it.get("dimension", "")).startswith("内容质量-数值与计算"):
            (fails if strict else notes).append(f"{rid} 含表格类定位语 {hit_words}，确认是否贴切")
        if str(it.get("criterion_type")) == "Subjective":
            notes.append(f"{rid} 标为 Subjective，确认能否按可核验锚点改 Objective")
        if DUP_SCALE.search(desc):
            (fails if strict else notes).append(
                f"{rid} description 里重复给出 0–1 档位段（档位唯一来源是 levels；"
                f"渲染到 toml 会变成两套标度，验收见 validate_rubrics.py 的"
                f"「toml 中无旧 0–1 标度残留」/「likert 锚点一致性」）")
        m2 = META_PATCH.search(desc)
        if m2:
            (fails if strict else notes).append(
                f"{rid} 判据内出现替题面兜底的解释语 {m2.group(0)!r}"
                f"（题面有歧义才会写这句，应回 instruction.md 改，判据只留判定标准）")
        anchors = set()
        for quoted in ANCHOR_QUOTED.findall(desc):
            if any(k in quoted for k in ("口径", "阈值", "限额", "标准")) \
                    and not any(s in quoted for s in QUOTED_SKIP):
                anchors.add(quoted)
        for a in sorted(anchors):
            if materials is None:
                notes.append(f"{rid} 锚点待人工核实：{a}")
            elif a not in materials:
                notes.append(f"{rid} 锚点 {a} 未在 input_files 命中，请确认出处"
                             f"（材料为英文原文时中文措辞不会命中，需人工判断）")
        if materials is not None:
            squashed = materials.replace(" ", "").replace("\n", "")
            # 本地增量：材料里的数字常有千分位/全角/换行差异，用纯数字串再比一次，压掉误报
            digits_only = re.sub(r"[^\d.]", "", materials)
            unverified = []
            for num in sorted(set(ANCHOR_NUMBER.findall(desc))):
                token = num.strip()
                if not token:
                    continue
                variants = {token, token.replace(" ", ""), token.replace(",", ""),
                            re.sub(r"[^\d.]", "", token)}
                if any(v and (v in squashed or v in digits_only) for v in variants):
                    continue
                unverified.append(token)
            if unverified:
                head = "、".join(unverified[:8])
                more = f"（另有 {len(unverified) - 8} 项）" if len(unverified) > 8 else ""
                notes.append(f"{rid} 数值/条款未在 input_files 直接命中，请人工确认：{head}{more}"
                             f"（派生计算值、单位换算或格式差异属正常，但自造口径不行）")
    for msg in fails:
        print("  FAIL " + msg)
    for msg in notes:
        print("  NOTE " + msg)
    if not fails and not notes:
        print("  OK   未发现风格问题")
    print(f"  —— 提问式/量词/措辞问题 {len(fails)} 条，待确认 {len(notes)} 条"
          f"（材料文本比对：{'已启用' if materials is not None else '未启用'}）")
    return fails, notes


def main():
    global SKIP_MATERIAL
    args = sys.argv[1:]
    strict = "--strict" in args
    SKIP_MATERIAL = "--no-material" in args
    dirs = [a for a in args if not a.startswith("-")]
    if not dirs:
        raise SystemExit(__doc__)
    total_fail = 0
    for task_dir in dirs:
        fails, _notes = style_check(task_dir.rstrip("/\\"), strict)
        total_fail += len(fails)
    print(f"\nFAIL 合计: {total_fail}"
          + ("" if strict else "（当前为提示模式；加 --strict 可把风格项计为 FAIL）"))
    return 1 if total_fail else 0


if __name__ == "__main__":
    sys.exit(main())
