#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""独立复算核验（QC §7.5）：只用 environment/input_files/ 重建关键锚点，
不复用 solution/golden_output 的任何代码或结论，金标仅用于事后对照。"""
from __future__ import annotations

import csv
import os
import sys
from openpyxl import load_workbook

IN = sys.argv[1] if len(sys.argv) > 1 else None
if not IN or not os.path.isdir(IN):
    BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    IN = os.path.join(BASE, "FIN3-WKN-152", "environment", "input_files")
    if not os.path.isdir(IN):
        print("input dir not found:", IN)
        sys.exit(2)


def rows(path, sheet=None):
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def kv(path, sheet):
    out = {}
    for r in rows(path, sheet)[1:]:
        if r and r[0]:
            out[str(r[0]).strip()] = r
    return out


def num(x):
    if x in (None, ""):
        return 0.0
    return float(str(x).replace(",", ""))


def csv_rows(path):
    with open(path, encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


# --- 独立取数 ---------------------------------------------------------------
fin = {}
for r in rows(os.path.join(IN, "sec_filings", "SEC-01_financials_extract.xlsx"),
              "Public_Financials")[1:]:
    fin[str(r[0]).strip()] = {"FY2022": num(r[1]), "FY2023": num(r[2])}

ua = kv(os.path.join(IN, "committee", "Underwriting_Assumptions_20240320.xlsx"), "Assumptions")
ot = kv(os.path.join(IN, "committee", "Offering_Terms_20240320.xlsx"), "Offering_Terms")

rev23 = fin["Revenue"]["FY2023"]
adj = fin["Adjusted EBITDA"]["FY2023"]
sbc = fin["Stock-based compensation & related taxes"]["FY2023"]
cash = fin["Cash & cash equivalents"]["FY2023"]
msec = fin["Marketable securities"]["FY2023"]
fcf = fin["Free Cash Flow"]["FY2023"]

growth = num(ua["2024E revenue growth"][1])
p_lo = num(ua["Peer low EV/Revenue"][1])
p_mid = num(ua["Peer midpoint EV/Revenue"][1])
p_hi = num(ua["Peer high EV/Revenue"][1])
disc = num(ua["IPO discount to peer-implied equity"][1])

price = num(ot["Proposed Committee Price"][1])
primary = num(ot["Primary shares offered"][1])
secondary = num(ot["Secondary shares offered"][1])
gs = num(ot["Greenshoe shares"][2])          # Full Greenshoe 列
pre_sh = num(ot["Pre-money economic shares"][1])
fee = num(ot["Underwriting fee assumption"][1])
fixed = num(ot["Fixed company offering expenses"][1])

# 月度明细：去重 + 排除未审计
fy = {"FY2022": 0.0, "FY2023": 0.0}
seen, dropped = set(), []
for r in csv_rows(os.path.join(IN, "financials", "monthly_revenue_2022_2023.csv")):
    st = r["status"]
    if st in ("subtotal", "duplicate-export", "unaudited-preliminary"):
        dropped.append((r["fy"], r["month"], st))
        continue
    k = (r["fy"], r["month"])
    if k in seen:
        dropped.append((r["fy"], r["month"], "dup-key"))
        continue
    seen.add(k)
    fy[r["fy"]] += num(r["revenue_usd_mm"])

# --- 独立计算 ---------------------------------------------------------------
uw_ebitda = adj - sbc
rev24e = rev23 * (1 + growth)
net_cash = cash + msec
eq_lo, eq_mid, eq_hi = (p_lo * rev24e + net_cash), (p_mid * rev24e + net_cash), (p_hi * rev24e + net_cash)
val_lo = eq_lo / pre_sh * (1 - disc)
val_mid = eq_mid / pre_sh * (1 - disc)
val_hi = eq_hi / pre_sh * (1 - disc)
gross_primary = primary * price
net_primary = gross_primary - gross_primary * fee - fixed
post = pre_sh + primary
full_post = post + gs
gs_incr = gs * price * (1 - fee)

CHECKS = [
    ("Revenue FY2023", rev23, 804.029, 0.005),
    ("Underwriting EBITDA", uw_ebitda, -118.361, 0.05),
    ("2024E Revenue", rev24e, 980.91538, 0.01),
    ("Net cash bridge", net_cash, 1213.122, 0.005),
    ("Offer low", val_lo, 31.27, 0.10),
    ("Offer mid", val_mid, 34.26, 0.10),
    ("Offer high", val_hi, 37.25, 0.10),
    ("Price vs midpoint", price - val_mid, -0.26, 0.05),
    ("Gross primary", gross_primary, 519.402, 0.5),
    ("Net primary", net_primary, 486.432, 0.5),
    ("Post-money shares", post, 158.993090, 0.001),
    ("Full-greenshoe shares", full_post, 162.293090, 0.001),
    ("Greenshoe incremental net", gs_incr, 106.590, 0.5),
    ("Monthly FY2022 tie-out", fy["FY2022"], 666.701, 0.005),
    ("Monthly FY2023 tie-out", fy["FY2023"], 804.029, 0.005),
]

print(f"input dir: {IN}")
print(f"{'anchor':28s}{'independent':>16s}{'golden':>16s}  verdict")
bad = 0
for name, got, exp, tol in CHECKS:
    ok = abs(got - exp) <= tol
    bad += 0 if ok else 1
    print(f"{name:28s}{got:16.6f}{exp:16.6f}  {'OK' if ok else 'MISMATCH'}")
print()
print("dropped detail rows:", dropped)
print("FCF FY2023 =", fcf, "| peer range", p_lo, p_mid, p_hi, "| discount", disc)
print("RESULT:", "ALL ANCHORS AGREE" if bad == 0 else f"{bad} MISMATCH(ES)")
sys.exit(0 if bad == 0 else 1)
