# -*- coding: utf-8 -*-
"""核查 102600003.IB 与华骏集团系在十个账户中的持仓分布与市值。"""
import csv
import io
import json
from pathlib import Path

import openpyxl

PKG = Path(r"harbor-weakness\work_fin-b01_20261012-153\FIN3-WKN-153")
INP = PKG / "environment" / "input_files"
ASOF = "2026-09-30"


def read_csv(p):
    with open(p, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


master = {r["sec_code"]: r for r in read_csv(INP / "portfolio" / "security_master.csv")}
prices = {}
for r in read_csv(INP / "data" / "price_feed_20260930_repush.csv"):
    if r["trade_date"] == ASOF:
        prices[r["instrument_id"]] = float(r["close_price"])

wb = openpyxl.load_workbook(INP / "portfolio" / "holdings_2026Q3.xlsx", read_only=True, data_only=True)
rows = []
for name in wb.sheetnames:
    if not name.startswith("acct_"):
        continue
    ws = wb[name]
    head = [str(c).strip() if c is not None else "" for c in next(ws.iter_rows(values_only=True))]
    idx = {h: i for i, h in enumerate(head)}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r is None or r[idx["sec_code"]] in (None, ""):
            continue
        rows.append({"sheet": name, "acct": str(r[idx["acct"]]), "sec_code": str(r[idx["sec_code"]]),
                     "vol": float(r[idx["vol"]]), "as_of": str(r[idx["as_of"]]),
                     "src": str(r[idx["source_row_id"]])})
wb.close()


def up(code):
    t = master[code]["sec_type"]
    p = prices.get(code, 0.0)
    return p if t == "EquityFund" else p / 100.0


# 口径 A：与金标一致（估值日 + active + 去重）
valid, seen = [], set()
for r in rows:
    sec = master.get(r["sec_code"])
    if r["as_of"] != ASOF or sec is None or sec["status"] != "active":
        continue
    key = (r["acct"], r["sec_code"], r["vol"], r["as_of"])
    if key in seen:
        continue
    seen.add(key)
    valid.append(r)

target = "102600003.IB"
print("=== 102600003.IB 全部记录（含未过滤） ===")
for r in rows:
    if r["sec_code"] == target:
        print(f"  {r['sheet']} vol={r['vol']:,.0f} as_of={r['as_of']} src={r['src']} "
              f"price={prices.get(target)}")
tv = sum(r["vol"] for r in valid if r["sec_code"] == target)
print(f"  有效记录合计面值 = {tv:,.0f}，市值 = {tv * up(target):,.2f}")

print("=== 华骏集团系全部有效记录 ===")
hj = [r for r in valid if master[r["sec_code"]]["issuer"] == "华骏集团"]
for code in sorted({r["sec_code"] for r in hj}):
    sub = [r for r in hj if r["sec_code"] == code]
    vol = sum(r["vol"] for r in sub)
    print(f"  {code} {master[code]['sec_name']} 账户 {[r['acct'] for r in sub]} "
          f"面值合计={vol:,.0f} 市值={vol * up(code):,.2f}")
hj_mv = sum(r["vol"] * up(r["sec_code"]) for r in hj)
print(f"  华骏系合计市值 = {hj_mv:,.2f}")

# 分母
nav = {r["item"]: r["value"] for r in read_csv(INP / "valuation" / "nav_report_20260930.csv")}
accr = read_csv(INP / "valuation" / "accruals_20260930.csv")
recv = sum(float(r["amount"]) for r in accr if r["item"].startswith(("interest", "dividend")))
pay = sum(float(r["amount"]) for r in accr if r["item"].startswith("fee_payable"))
cash = float(nav["cash_and_equivalent"])
mv_all = sum(r["vol"] * up(r["sec_code"]) for r in valid)
T = mv_all + cash + recv - pay
print(f"=== 分母 资产净值 = {T:,.2f} ===")

# 单一证券（跨账户合计）最大值
by_code = {}
for r in valid:
    by_code[r["sec_code"]] = by_code.get(r["sec_code"], 0.0) + r["vol"] * up(r["sec_code"])
top = sorted(by_code.items(), key=lambda kv: -kv[1])[:5]
print("=== 按证券跨账户合计市值 Top5 ===")
for code, mv in top:
    print(f"  {code} {master[code]['sec_name'][:24]} 市值={mv:,.2f} 占比={mv / T * 100:.4f}%")

# 单一发行人
by_issuer = {}
for r in valid:
    iss = master[r["sec_code"]]["issuer"]
    by_issuer[iss] = by_issuer.get(iss, 0.0) + r["vol"] * up(r["sec_code"])
ti = sorted(by_issuer.items(), key=lambda kv: -kv[1])[:3]
print("=== 按发行人合计 Top3 ===")
for iss, mv in ti:
    print(f"  {iss} 市值={mv:,.2f} 占比={mv / T * 100:.4f}%")

# 单账户单券（金标旧口径）
row_max = max(valid, key=lambda r: r["vol"] * up(r["sec_code"]))
print(f"=== 金标旧口径（单账户单行）最大 = {row_max['sec_code']} {row_max['acct']} "
      f"市值={row_max['vol'] * up(row_max['sec_code']):,.2f} "
      f"占比={row_max['vol'] * up(row_max['sec_code']) / T * 100:.4f}% ===")

anchors = json.load(io.open(r"harbor-weakness\_qc_runs\tools\anchors_153.json", encoding="utf-8"))
print("=== anchors 记录 ===")
print("  max_sec:", anchors["max_sec"])
print("  huajun_ratio_pct:", anchors["huajun_ratio_pct"])
