# -*- coding: utf-8 -*-
"""独立抽验新金标：网格/Tieout_Detail/算式/清点/理由/核验四条/分列。
（与子代理自检互为交叉验证）"""
import pathlib
import re
import sys

sys.stdout.reconfigure(encoding="utf-8")
from openpyxl import load_workbook

TASK = pathlib.Path(r"C:\Users\Administrator\Desktop\wff-task\harbor-weakness\FIN3-WKN-152")
G = TASK / "solution" / "golden_output"

print("=" * 96)
print("1) R18 敏感性矩阵网格（须恰为 18/20/22/24/26 × 4.0/4.3/4.5/4.8/5.0）")
print("=" * 96)
vm = (G / "FIN3-WKN-152_valuation_matrix.csv").read_text(encoding="utf-8", errors="replace")
lines = [l for l in vm.splitlines() if l.strip()]
for i, l in enumerate(lines[:12]):
    print(f"  {i}: {l[:150]}")
print(f"  共 {len(lines)} 行")

print()
print("=" * 96)
print("2) R35 Tieout_Detail 工作表")
print("=" * 96)
wb = load_workbook(G / "FIN3-WKN-152_ipo_model.xlsx", read_only=True, data_only=True)
print(f"  工作表({len(wb.sheetnames)}): {wb.sheetnames}")
if "Tieout_Detail" in wb.sheetnames:
    ws = wb["Tieout_Detail"]
    rows = [r for r in ws.iter_rows(values_only=True) if any(x is not None for x in r)]
    print(f"  行数 {len(rows)}（表头 + 数据）")
    for r in rows[:9]:
        print(f"    {str(r)[:160]}")
else:
    print("  [!!] 无 Tieout_Detail")
wb.close()

print()
print("=" * 96)
print("3) R34 材料清点（60 文件 / 7 目录 + 分项）")
print("=" * 96)
memo = (G / "FIN3-WKN-152_pricing_memo.md").read_text(encoding="utf-8", errors="replace")
for pat in [r"60", r"committee", r"sec_filings", r"financials", r"comps", r"research",
            r"internal", r"legacy"]:
    hit = [l.strip()[:150] for l in memo.splitlines() if pat in l]
    if hit and pat in ("60", "committee"):
        print(f"  [{pat}] {hit[:2]}")
# 找目录分项行
seg = [l for l in memo.splitlines() if re.search(r"(8|17|14|4|5|6|4).*(文件|files)", l)]
print(f"  分项计数行: {seg[:4]}")

print()
print("=" * 96)
print("4) R05/R07/R08/R11/R17 关键算式是否在交付物中")
print("=" * 96)
need = {
    "R05 EBITDA": ["-118.361", "-69.275", "49.086"],
    "R07 2024E": ["980.915", "804.029", "22%"],
    "R08 净现金桥": ["1,213.122", "401.176", "811.946", "1213.122"],
    "R11 募集费用": ["519.402", "486.432", "25.970", "7.000"],
    "R17 股本桥": ["158.993090", "143.716563", "15.276527", "162.293090"],
}
xlsx_txt = ""
wb2 = load_workbook(G / "FIN3-WKN-152_ipo_model.xlsx", read_only=True, data_only=True)
for sn in wb2.sheetnames:
    for row in wb2[sn].iter_rows(values_only=True):
        xlsx_txt += " ".join(str(c) for c in row if c is not None) + "\n"
wb2.close()
qoe = (G / "FIN3-WKN-152_qoe_bridge.csv").read_text(encoding="utf-8", errors="replace")
st = (G / "FIN3-WKN-152_source_trace.csv").read_text(encoding="utf-8", errors="replace")
corpus = {"memo": memo, "xlsx": xlsx_txt, "qoe": qoe, "source_trace": st}
for label, toks in need.items():
    where = [name for name, txt in corpus.items() if all(t in txt for t in toks)]
    print(f"  {label:<16} 完整出现在: {where or '（未找到全部）'}")
    for name, txt in corpus.items():
        miss = [t for t in toks if t not in txt]
        if name in ("memo", "xlsx") and miss and not where:
            print(f"      {name} 缺: {miss}")

print()
print("=" * 96)
print("5) R33 数据核验四条分列 / R36 四项理由 / R06 两项负前提")
print("=" * 96)
checks = {
    "核验① 2022-08 重复": "2022-08" in memo,
    "核验② 2023-12 双值": "2023-12" in memo,
    "核验③ 分部量级": ("152.470" in memo or "15.247" in memo),
    "核验④ SBC TOTAL": ("49.680" in memo or "49.086" in memo),
    "R36① 中点4.5x理由": ("4.5" in memo and ("中点" in memo or "midpoint" in memo.lower())),
    "R36② flash 排除": "flash" in memo.lower(),
    "R36③ v3 非 v2": ("v3" in memo or "v2" in memo),
    "R36④ primary 费基": "primary" in memo.lower(),
    "R06 负前提 EBITDA": ("-118.361" in memo or "EBITDA" in memo),
    "R06 负前提 FCF": "FCF" in memo,
}
for k, v in checks.items():
    print(f"  [{'OK' if v else '!!'}] {k}")

print()
print("=" * 96)
print("6) R02 备忘录字数（≤1600 中文字）")
print("=" * 96)
n = len(re.findall(r"[一-鿿]", memo))
print(f"  中文字数 {n} / 1600  {'OK' if n <= 1600 else '!! 超限'}")

print()
print("=" * 96)
print("7) source_trace 竞争值（bank comps 4.2x-5.1x + v2 3.5x-5.5x）与 Priority")
print("=" * 96)
for k in ["4.2", "5.1", "3.5x", "5.5", "Priority", "P1", "P9"]:
    print(f"  [{'OK' if k in st else '!!'}] {k} 在 source_trace: {k in st}")
