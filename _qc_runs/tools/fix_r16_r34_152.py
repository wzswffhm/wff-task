# -*- coding: utf-8 -*-
"""修复两条判据自身的矛盾（子代理如实上报，不伪造）：
1) R16 结果数 632.456518 -> 631.601918（真值；原值与自身算式差 0.8546 > ±0.5 容差）
2) R34 七目录小计 58，补「根目录散件 2」项，使 58 + 2 = 60 自洽
双文件同步（toml + json）。
"""
import json
import pathlib
import sys

import tomllib

sys.stdout.reconfigure(encoding="utf-8")

TASK = pathlib.Path(r"C:\Users\Administrator\Desktop\wff-task\harbor-weakness\FIN3-WKN-152")
TOML = TASK / "tests" / "rubrics.toml"
JSONF = TASK / "rubrics.json"

FIX = {
"R16": (
    "greenshoe 处理正确：**3.3m** greenshoe 不预先并入 Base 情景，仅单列 full-exercise 情景；"
    "full exercise 下公司 gross primary proceeds 为 (15.276527 + 3.3)m × $34 = **631.601918mm**"
    "（允许 ±0.5mm），并须写出该算式（列报值齐全）。把 greenshoe 预先并入 Base、"
    "或未单列 full-exercise 情景、或未写出该算式即不满足。 "
    "Deliverables to inspect: `/app/output/FIN3-WKN-152_ipo_model.xlsx`."
),
"R34": (
    "输入材料清点正确：交付物中明确给出 `/app/input_files/` 的**文件总数 60** 与**来源目录数 7**，"
    "并**分列七个来源目录的文件数**（`committee/` 8、`sec_filings/` 17、`financials/` 14、"
    "`comps/` 4、`research/` 5、`internal/` 6、`legacy/` 4）以及**根目录散件数 2**"
    "（`00_README.md`、`data_dictionary.md`），使**七目录小计 58 + 根目录散件 2 = 总数 60**，"
    "即分项合计与总数勾稽（差 1 即不满足）。总数或目录数错误、缺任一目录的分项计数、"
    "缺根目录散件项、分项合计与总数不符、或只给总数未分列目录，即不满足。 "
    "Deliverables to inspect: `/app/output/FIN3-WKN-152_ipo_model.xlsx`, "
    "`/app/output/FIN3-WKN-152_pricing_memo.md`."
),
}

# ── toml ──
raw = TOML.read_text(encoding="utf-8")
t = tomllib.loads(raw)
for cid, new in FIX.items():
    old = next(c["description"] for c in t["criterion"] if c["id"] == cid)
    if old not in raw:
        print(f"  [!!] {cid} 原文未在 toml 定位")
        sys.exit(1)
    if old == new:
        print(f"  [--] {cid} 已是目标文本")
        continue
    raw = raw.replace(old, new, 1)
    print(f"  [OK] toml {cid}: 替换（{len(old)} -> {len(new)} 字）")
TOML.write_text(raw, encoding="utf-8", newline="\n")

# ── json ──
j = json.loads(JSONF.read_text(encoding="utf-8"))
for it in j["items"]:
    if it["id"] in FIX:
        it["description"] = FIX[it["id"]]
JSONF.write_text(json.dumps(j, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")
print("  [OK] json 同步 2 条")

# ── 校验 ──
t2 = tomllib.loads(TOML.read_text(encoding="utf-8"))
j2 = json.loads(JSONF.read_text(encoding="utf-8"))
jt = {c["id"]: c for c in t2["criterion"]}
jj = {i["id"]: i for i in j2["items"]}
print()
print("  ID 一致:", set(jt) == set(jj), f"({len(jt)} 条)")
bad = [k for k in jt if jt[k]["description"] != jj[k]["description"]]
print("  description 双文件一致:", not bad, bad if bad else "")
pos = sum(float(c["weight"]) for c in t2["criterion"] if not c.get("negate"))
print(f"  正分池 {pos} / s_max {j2['metadata']['scoring']['s_max']}  "
      f"一致 {pos == j2['metadata']['scoring']['s_max']}")
print()
print("  R16 尾:", jt["R16"]["description"][:150], "…")
print("  R34 中:", jt["R34"]["description"][60:230], "…")

# 与金标实际输出交叉核对
from openpyxl import load_workbook
G = TASK / "solution" / "golden_output"
memo = (G / "FIN3-WKN-152_pricing_memo.md").read_text(encoding="utf-8", errors="replace")
wb = load_workbook(G / "FIN3-WKN-152_ipo_model.xlsx", read_only=True, data_only=True)
xt = "\n".join(" ".join(str(c) for c in r if c is not None)
               for sn in wb.sheetnames for r in wb[sn].iter_rows(values_only=True))
wb.close()
print()
print("=" * 92)
print("与金标交叉核对（新判据要求 vs 金标实际输出）")
print("=" * 92)
print(f"  R16 631.601918:  memo={('631.601918' in memo)}  xlsx={('631.601918' in xt)}")
print(f"  R34 散件2:       memo={('散件' in memo and '2' in memo)}")
print(f"  R34 58+2=60:     memo={('58' in memo)}")
print(f"  R34 合计60:      memo={('60' in memo)}")
